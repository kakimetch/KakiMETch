from collections.abc import Iterator, Mapping
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

import openpyxl

from app.database import get_connection

DATA_DIRECTORY = Path(__file__).resolve().parents[2] / "data"
WEEKDAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]


def parse_excel_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            return datetime.strptime(str(int(value)), "%Y%m%d").date()
        except ValueError:
            return None
    return None


def mobility_from_equipment(wheelchair: object, walking_frame: object) -> str:
    """Derive the initial imported mobility value from the available source fields."""
    if is_yes(wheelchair):
        return "wheelchair_user"
    if is_yes(walking_frame):
        return "walking_frame_user"
    return "ambulant"


def expand_available_days(value: object) -> list[str]:
    days = as_text(value).split("/")
    if len(days) == 1 and "-" in days[0]:
        start_day, end_day = (part.strip().title() for part in days[0].split("-", 1))
        if start_day in WEEKDAYS and end_day in WEEKDAYS:
            return WEEKDAYS[WEEKDAYS.index(start_day) : WEEKDAYS.index(end_day) + 1]

    return [day.strip().title() for day in days if day.strip().title() in WEEKDAYS]


def load_demo_data(data_directory: Path = DATA_DIRECTORY) -> tuple[int, int, int]:
    """Import the provided demo workbooks and return client, escort, and trip counts."""
    with get_connection() as connection:
        with connection.cursor() as cursor:
            client_count = load_clients(
                cursor, data_directory / "Dummy_MasterData_Updated.xlsx"
            )
            escort_count = load_escorts(
                cursor, data_directory / "Dummy_EscortRoster.xlsx"
            )
            trip_count = load_trips(
                cursor, data_directory / "Dummy_ScheduleData_September_2026.xlsx"
            )

    return client_count, escort_count, trip_count


def load_clients(cursor: Any, workbook_path: Path) -> int:
    count = 0
    for row in read_rows(workbook_path):
        mobility_status = mobility_from_equipment(
            row["Wheelchair (WC)"], row["Walking Frame/ Stick"]
        )
        cursor.execute(
            """
            insert into public.elderly_clients (
                name, nric, aic_registration_no, postal_code, block, unit, street_name,
                address, contact_no, caregiver_name, escort_required, co_payment,
                date_of_birth, nmts_effective_date, nmts_expired_date, date_of_entry,
                action_updated_date, lh_service_agreement, sw_service_agreement,
                wheelchair_required, walking_frame_required, caregiver_or_maid_available,
                gender, address_source, dialect, weight_kg, aic_mobility_status,
                lh_mobility_status
            ) values (
                %(name)s, %(nric)s, %(aic_registration_no)s, %(postal_code)s, %(block)s,
                %(unit)s, %(street_name)s, %(address)s, %(contact_no)s, %(caregiver_name)s,
                %(escort_required)s, %(co_payment)s, %(date_of_birth)s,
                %(nmts_effective_date)s, %(nmts_expired_date)s, %(date_of_entry)s,
                %(action_updated_date)s, %(lh_service_agreement)s, %(sw_service_agreement)s,
                %(wheelchair_required)s, %(walking_frame_required)s,
                %(caregiver_or_maid_available)s, %(gender)s, %(address_source)s, %(dialect)s,
                %(weight_kg)s, %(aic_mobility_status)s, %(lh_mobility_status)s
            ) on conflict (nric) do update set
                name = excluded.name,
                aic_registration_no = excluded.aic_registration_no,
                postal_code = excluded.postal_code,
                block = excluded.block,
                unit = excluded.unit,
                street_name = excluded.street_name,
                address = excluded.address,
                contact_no = excluded.contact_no,
                caregiver_name = excluded.caregiver_name,
                escort_required = excluded.escort_required,
                co_payment = excluded.co_payment,
                date_of_birth = excluded.date_of_birth,
                nmts_effective_date = excluded.nmts_effective_date,
                nmts_expired_date = excluded.nmts_expired_date,
                date_of_entry = excluded.date_of_entry,
                action_updated_date = excluded.action_updated_date,
                lh_service_agreement = excluded.lh_service_agreement,
                sw_service_agreement = excluded.sw_service_agreement,
                wheelchair_required = excluded.wheelchair_required,
                walking_frame_required = excluded.walking_frame_required,
                caregiver_or_maid_available = excluded.caregiver_or_maid_available,
                gender = excluded.gender,
                address_source = excluded.address_source,
                dialect = excluded.dialect,
                weight_kg = excluded.weight_kg,
                aic_mobility_status = excluded.aic_mobility_status,
                lh_mobility_status = excluded.lh_mobility_status,
                updated_at = now()
            """,
            {
                "name": as_text(row["NAME"]),
                "nric": as_text(row["NRIC"]),
                "aic_registration_no": as_text(row["AIC Reg No"]),
                "postal_code": as_text(row["Postal Code"]),
                "block": as_text(row["BLK"]),
                "unit": as_text(row["UNIT"]),
                "street_name": as_text(row["St Name"]),
                "address": build_address(row),
                "contact_no": as_text(row["Contact No"]),
                "caregiver_name": as_text(row["Caregiver"]),
                "escort_required": is_yes(row["Escort (Y/N)"]),
                "co_payment": row["Co-payment"],
                "date_of_birth": parse_excel_date(row["DOB (YYYYMMDD)"]),
                "nmts_effective_date": parse_excel_date(
                    row["NMTS effective date (YYYYMMDD)"]
                ),
                "nmts_expired_date": parse_excel_date(
                    row["NMTS expired date (YYYYMMDD)"]
                ),
                "date_of_entry": parse_excel_date(row["Date Of Entry (YYYMMDD)"]),
                "action_updated_date": as_text(row["Action/updated date"]),
                "lh_service_agreement": as_text(row["LH Service Agreement"]) or None,
                "sw_service_agreement": as_text(row["SW Service Agreement"]) or None,
                "wheelchair_required": is_yes(row["Wheelchair (WC)"]),
                "walking_frame_required": is_yes(row["Walking Frame/ Stick"]),
                "caregiver_or_maid_available": is_yes(row["Caregiver/Maid"]),
                "gender": as_text(row["M/F"]) or None,
                "address_source": as_text(row["Address Source"]),
                "dialect": as_text(row["Dialect"]),
                "weight_kg": row["Weight (kg)"],
                "aic_mobility_status": mobility_status,
                "lh_mobility_status": mobility_status,
            },
        )
        count += 1

    return count


def load_escorts(cursor: Any, workbook_path: Path) -> int:
    count = 0
    for row in read_rows(workbook_path):
        cursor.execute(
            """
            insert into public.escorts (
                external_id, name, gender, dialects, available_days, available_timeslot,
                wheelchair_handling_capable, contact_no
            ) values (
                %(external_id)s, %(name)s, %(gender)s, %(dialects)s, %(available_days)s,
                %(available_timeslot)s, %(wheelchair_handling_capable)s, %(contact_no)s
            ) on conflict (external_id) do update set
                name = excluded.name,
                gender = excluded.gender,
                dialects = excluded.dialects,
                available_days = excluded.available_days,
                available_timeslot = excluded.available_timeslot,
                wheelchair_handling_capable = excluded.wheelchair_handling_capable,
                contact_no = excluded.contact_no,
                updated_at = now()
            """,
            {
                "external_id": as_text(row["Escort ID"]),
                "name": as_text(row["Name"]),
                "gender": as_text(row["Gender (M/F)"]),
                "dialects": split_dialects(row["Dialect(s) Spoken"]),
                "available_days": expand_available_days(row["Available Days"]),
                "available_timeslot": as_text(row["Available Timeslot"]),
                "wheelchair_handling_capable": is_yes(row["Wheelchair Handling (Y/N)"]),
                "contact_no": as_text(row["Contact No"]),
            },
        )
        count += 1

    return count


def load_trips(cursor: Any, workbook_path: Path) -> int:
    count = 0
    for row in read_rows(workbook_path):
        appt_date = parse_excel_date(row["Appt Date"])
        destination = as_text(row["Destination"])
        if appt_date is None or not destination:
            continue

        for appointment_time in (row["1st Appt"], row["2nd Appt"]):
            if not isinstance(appointment_time, time):
                continue
            cursor.execute(
                """
                insert into public.trips (elderly_id, appt_date, appt_time, destination)
                select elderly_clients.id, %s, %s, %s
                from public.elderly_clients
                where elderly_clients.nric = %s
                  and not exists (
                      select 1
                      from public.trips
                      where trips.elderly_id = elderly_clients.id
                        and trips.appt_date = %s
                        and trips.appt_time = %s
                        and trips.destination = %s
                  )
                """,
                (
                    appt_date,
                    appointment_time,
                    destination,
                    as_text(row["NRIC"]),
                    appt_date,
                    appointment_time,
                    destination,
                ),
            )
            count += cursor.rowcount

    return count


def read_rows(workbook_path: Path) -> Iterator[dict[str, object]]:
    worksheet = openpyxl.load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
    ).active
    headers = next(worksheet.iter_rows(values_only=True))
    for values in worksheet.iter_rows(min_row=2, values_only=True):
        yield dict(zip(headers, values, strict=True))


def build_address(row: Mapping[str, object]) -> str:
    return " ".join(
        part
        for part in (as_text(row["BLK"]), as_text(row["St Name"]), as_text(row["UNIT"]))
        if part
    )


def split_dialects(value: object) -> list[str]:
    return [dialect.strip() for dialect in as_text(value).split(",") if dialect.strip()]


def is_yes(value: object) -> bool:
    return as_text(value).upper() == "Y"


def as_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


if __name__ == "__main__":
    clients, escorts, trips = load_demo_data()
    print(f"Imported {clients} clients, {escorts} escorts, and {trips} new trips.")
